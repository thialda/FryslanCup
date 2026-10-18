# -*- coding: utf-8 -*-
# Tussenstand Friese sloepen: som van de 3 beste wedstrijden (minste punten).
# Sloepen met minder dan 3 gefinishte wedstrijden vallen buiten het klassement.
import os, sys
sys.path.insert(0, os.path.dirname(__file__))
from make_friese_uitslagen import (HKS_HEREN, GROU_HEREN, LEMMER_HEREN, MAKKUM_HEREN, LANGWEER_HEREN,
    HKS_DAMES, GROU_DAMES, LEMMER_DAMES, MAKKUM_DAMES, LANGWEER_DAMES, FRIES, frisian_ranking, OUT_DIR)

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BESTE = 3
KLASSEN = {
    "Heren": [("HKS", HKS_HEREN), ("Grou", GROU_HEREN), ("IFK Lemmer", LEMMER_HEREN),
              ("Slag om Makkum", MAKKUM_HEREN), ("Langweer", LANGWEER_HEREN)],
    "Dames": [("HKS", HKS_DAMES), ("Grou", GROU_DAMES), ("IFK Lemmer", LEMMER_DAMES),
              ("Slag om Makkum", MAKKUM_DAMES), ("Langweer", LANGWEER_DAMES)],
}

def fmt(v):
    if v is None:
        return "-"
    return "0,9" if v == 0.9 else str(int(v))

def fmt_tot(v):
    return f"{v:.1f}".replace(".", ",").replace(",0", "")

def bereken(race_cols):
    per = {}
    for naam, rows in race_cols:
        for sloep, punten in frisian_ranking(rows):
            per.setdefault(sloep, {})[naam] = punten
    deel, buiten = [], []
    for sloep, d in per.items():
        waarden = [d.get(n) for n, _ in race_cols]
        behaald = sorted(v for v in waarden if v is not None)
        verenig, plaats = FRIES[sloep]
        if len(behaald) < BESTE:
            buiten.append((sloep, len(behaald)))
            continue
        tel = behaald[:BESTE]
        # markeer de wedstrijden die meetellen (bij gelijke waarden de eerste keren)
        mee, rest = [], list(tel)
        for v in waarden:
            if v is not None and v in rest:
                rest.remove(v); mee.append(True)
            else:
                mee.append(False)
        deel.append((sloep, verenig, plaats, waarden, mee, sum(tel)))
    deel.sort(key=lambda x: (x[5], x[0]))
    return deel, buiten

styles = getSampleStyleSheet()
title_style = ParagraphStyle("t", parent=styles["Heading1"], fontSize=14, spaceAfter=2)
sub_style = ParagraphStyle("s", parent=styles["Normal"], fontSize=9, spaceAfter=1)
note_style = ParagraphStyle("n", parent=styles["Normal"], fontSize=7, textColor=colors.grey, spaceBefore=8)

for klasse, race_cols in KLASSEN.items():
    deel, buiten = bereken(race_cols)
    namen = [n for n, _ in race_cols]
    uitleg = (f"Som van de {BESTE} beste wedstrijden (minste punten, 1e = 0,9 punt, 2e = 2 punten, elke volgende plaats 1 punt meer in het "
              f"Friese klassement). Een slechtere uitslag valt weg (grijs, tussen haakjes). Sloepen met minder dan {BESTE} gefinishte "
              f"wedstrijden staan niet in het klassement ({len(buiten)} sloepen). Stand na {len(race_cols)} wedstrijden.")
    # PDF
    doc = SimpleDocTemplate(os.path.join(OUT_DIR, f"Tussenstand beste {BESTE} - Friese {klasse.lower()}sloepen.pdf"),
        pagesize=landscape(A4), leftMargin=15*mm, rightMargin=15*mm, topMargin=15*mm, bottomMargin=15*mm)
    data = [["#", "Sloep", "Vereniging", "Plaats"] + namen + [f"Beste {BESTE}"]]
    for i, (sloep, ver, pl, waarden, mee, tot) in enumerate(deel, 1):
        cells = [fmt(v) if (v is None or m) else f"({fmt(v)})" for v, m in zip(waarden, mee)]
        data.append([str(i), sloep, ver, pl] + cells + [fmt_tot(tot)])
    t = Table(data, colWidths=[8*mm, 34*mm, 55*mm, 30*mm] + [22*mm]*len(namen) + [20*mm], repeatRows=1)
    st = [("FONTSIZE", (0,0), (-1,-1), 8), ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
          ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#dbe9f4")), ("GRID", (0,0), (-1,-1), 0.4, colors.grey),
          ("ALIGN", (0,0), (0,-1), "CENTER"), ("ALIGN", (4,0), (-1,-1), "CENTER"),
          ("FONTNAME", (-1,1), (-1,-1), "Helvetica-Bold"),
          ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#f4f7fa")])]
    for r, (sloep, *_rest) in enumerate(deel, 1):
        for c, m in enumerate(deel[r-1][4]):
            if not m:
                st.append(("TEXTCOLOR", (4+c, r), (4+c, r), colors.grey))
    t.setStyle(TableStyle(st))
    doc.build([Paragraph(f"Tussenstand Friese {klasse.lower()}sloepen 2026 — beste {BESTE} wedstrijden", title_style),
               Paragraph(uitleg, sub_style), Spacer(1, 4*mm), t,
               Paragraph("Niet officieel; Friese selectie samengesteld op basis van sloeproeien.nl en opgave van de gebruiker.", note_style)])
    # Excel
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = f"Tussenstand {klasse}"
    ws.append(["#", "Sloep", "Vereniging", "Plaats"] + namen + [f"Beste {BESTE}"])
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", start_color="DBE9F4")
    for i, (sloep, ver, pl, waarden, mee, tot) in enumerate(deel, 1):
        ws.append([i, sloep, ver, pl] + [v if (v is None or m) else f"({fmt(v)})" for v, m in zip(waarden, mee)] + [tot])
    for i, w in enumerate([5, 24, 36, 20] + [14]*len(namen) + [10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.append([]); ws.append([uitleg])
    wb.save(os.path.join(OUT_DIR, f"Tussenstand beste {BESTE} - Friese {klasse.lower()}sloepen.xlsx"))
    print(klasse, len(deel), "in klassement,", len(buiten), "erbuiten")
    for i, (sloep, ver, pl, waarden, mee, tot) in enumerate(deel[:10], 1):
        print(i, sloep, [fmt(v) for v in waarden], fmt_tot(tot))
    for i, x in enumerate(deel, 1):
        if x[0] == "Twirre": print("Twirre plek", i, [fmt(v) for v in x[3]], fmt_tot(x[5]))
